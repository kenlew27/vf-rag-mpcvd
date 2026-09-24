"""
fix_figure_placement_v2.py

Strategy: Remove ALL existing image paragraphs from manuscript_final.docx,
then re-insert each figure image directly above its own caption paragraph.
This guarantees correct figure-caption adjacency regardless of where images 
were placed by the original script.

Also applies all text edits from apply_all_edits.py.
"""

from __future__ import annotations
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

HERE = Path(__file__).parent
# Place manuscript_final.docx and figures/ in the repo root to use this script.
SRC  = HERE / "manuscript_final.docx"
DST  = HERE / "manuscript_final_fixed.docx"
FIGURES = HERE / "figures"

FIGURE_MAP = {
    "Fig. 1": FIGURES / "figure_1.png",
    "Fig. 2": FIGURES / "figure_2.png",
    "Fig. 3": FIGURES / "figure_3.png",
    "Fig. 4": FIGURES / "figure_4.png",
    "Fig. 5": FIGURES / "figure_5.png",
    "Fig. 6": FIGURES / "figure_6.jpg",
}

def has_picture(para) -> bool:
    return any(
        el.tag.endswith('}drawing') or el.tag.endswith('}pict')
        for el in para._p.iter()
    )

def para_starts_with_fig_caption(para, label: str) -> bool:
    return para.text.strip().startswith(label + " ")

def replace_in_para(para, old: str, new: str) -> bool:
    full = para.text
    if old not in full:
        return False
    new_full = full.replace(old, new)
    for run in para.runs:
        run.text = ""
    if para.runs:
        para.runs[0].text = new_full
    else:
        para.add_run(new_full)
    return True


# ── Text edits ────────────────────────────────────────────────────────────────

ABSTRACT_OLD = (
    "verification reduces the blocking-violation rate from 77% to 10% (87% relative reduction) "
    "while retry recovers delivery to approximately 90%. Two domain scientists independently rate "
    "92% of 50 end-to-end outputs as adequately grounded (score at least 4/5) and 86% as correct, "
    "with groundedness averaging 4.76/5. Stage-decomposed evaluation localises failures to specific "
    "pipeline components without conflating source-grounded verification with physical scientific "
    "validation."
)

ABSTRACT_NEW = (
    "verification reduces the blocking-violation rate \u2014 the fraction of outputs containing at "
    "least one unsupported or contradicted material proposition \u2014 from 77% to 10% (87% relative "
    "reduction), while retry recovers delivery (the fraction of outputs returned to the user) to "
    "90%. Two domain scientists independently rate 92% of 50 end-to-end outputs as adequately "
    "grounded (groundedness \u2265 4/5; every claim traceable to cited evidence) and 86% as correct "
    "(correctness \u2265 4/5; factually accurate), with mean groundedness 4.76/5. Scoring additionally "
    "captures scientific rigor (calibrated and appropriately bounded language), which averages "
    "4.27/5 with 89% of outputs rated \u2265 4/5. Stage-decomposed evaluation localises failures to "
    "specific pipeline components without conflating source-grounded verification with physical "
    "scientific validation."
)

CLAUDE_OLD = "Claude Sonnet 4 (June 2026)"
CLAUDE_NEW = "Claude Sonnet 4"

VOYAGE_OLD = "Voyage 4 embeddings (1,024-dimensional)"
VOYAGE_NEW = "Voyage 4 embeddings"

PORTABILITY_OLD = (
    "Cross-domain portability is therefore an architectural hypothesis, not an empirical result."
)
PORTABILITY_NEW = (
    "Whether the architecture works equally well in other domains remains to be tested."
)

DA_OLD_SNIPPET = "(iii) the blinded HTML grading workbench"
DA_OLD = (
    "The repository also includes: (i) the CVD diamond question set (50 questions across three "
    "complexity buckets); (ii) the evaluation rubric definitions for groundedness, correctness, "
    "and scientific rigor; (iii) the blinded HTML grading workbench used in the scientist "
    "evaluation; (iv) anonymised scientist evaluation scores (100 ratings from two independent "
    "graders); (v) scoring scripts used to compute all reported statistics, Wilson confidence "
    "intervals, and inter-rater agreement; and (vi) claim-verification benchmark fixture "
    "specifications (500 atomic claims with gold labels)."
)
DA_NEW = (
    "The repository also includes: (i) the CVD diamond question set (50 questions across three "
    "complexity buckets); (ii) the evaluation rubric definitions for groundedness, correctness, "
    "and scientific rigor; (iii) anonymised scientist evaluation scores (100 ratings from two "
    "independent graders); (iv) scoring scripts used to compute all reported statistics, Wilson "
    "confidence intervals, and inter-rater agreement; and (v) claim-verification benchmark "
    "fixture specifications (500 atomic claims with gold labels). The blinded scientist-evaluation "
    "grading workbench is not included in the public repository due to inclusion of session data."
)

AC_PLACEHOLDER = "Author contributions will be described using the CRediT taxonomy prior to publication."
AC_FULL = (
    "Author contributions are described using the CRediT (Contributor Roles Taxonomy) "
    "framework (https://credit.niso.org/). "
    "Bonnie Pang: Conceptualisation, Methodology, Software, Validation, Formal Analysis, "
    "Investigation, Writing \u2013 Review & Editing. "
    "Enzo Nakornsri: Conceptualisation, Methodology, Software, Validation, Formal Analysis, "
    "Investigation, Writing \u2013 Review & Editing. "
    "Ken Lew: Conceptualisation, Methodology, Software, Project Administration, Supervision, "
    "Writing \u2013 Original Draft, Writing \u2013 Review & Editing. "
    "Cristian Herrera Rodriguez: Investigation (scientist evaluation), Validation, "
    "Writing \u2013 Review & Editing. "
    "Thanh Tran: Investigation (scientist evaluation), Validation, Writing \u2013 Review & Editing."
)

GITHUB_PLACEHOLDER = "[REPOSITORY URL]"
GITHUB_URL = "https://github.com/kenlew27/vf-rag-mpcvd"

VERIFIER_INTRO_PARA = (
    "The claim verifier (Stage 5) operates as follows. After synthesis, deterministic markdown "
    "slicing extracts semantically coherent claim units from the generated answer. Each claim "
    "unit, together with its inline-cited evidence items (drawn from the typed evidence packet), "
    "is submitted as an isolated parcel to Claude Haiku 4.5. The verifier decomposes each claim "
    "into its constituent material propositions and assigns one of three statuses: supported "
    "(all propositions backed by cited evidence), unsupported (evidence absent, incomplete, "
    "irrelevant, or ambiguous), or contradicted (cited evidence explicitly disagrees). Universal "
    "terms and causal assertions receive additional scrutiny \u2014 one observation cannot establish "
    "a universal claim, and correlation is not causation. A single contradicted proposition makes "
    "the whole claim contradicted; a claim is supported only when all propositions are supported. "
    "Supported claims map deterministically to allow; unsupported and contradicted claims map to "
    "intervene. The verifier does not have access to information outside the cited evidence "
    "packet: it cannot search the web, query the database, or draw on its parametric knowledge "
    "to fill evidential gaps."
)

SECTION_42_TEXT = "4.2 Intervention and retry"


def apply_text_edits(doc: Document) -> dict:
    report = {}

    done = False
    for p in doc.paragraphs:
        if ABSTRACT_OLD in p.text:
            replace_in_para(p, ABSTRACT_OLD, ABSTRACT_NEW)
            done = True
            break
    report["Abstract"] = "OK" if done else "SKIP (not found / already edited)"

    count = 0
    for p in doc.paragraphs:
        if CLAUDE_OLD in p.text:
            replace_in_para(p, CLAUDE_OLD, CLAUDE_NEW)
            count += 1
    report["Claude (June 2026) removed"] = f"OK ({count})"

    count = 0
    for p in doc.paragraphs:
        if VOYAGE_OLD in p.text:
            replace_in_para(p, VOYAGE_OLD, VOYAGE_NEW)
            count += 1
    report["Voyage (1,024-dimensional) removed"] = f"OK ({count})"

    done = False
    for p in doc.paragraphs:
        if PORTABILITY_OLD in p.text:
            replace_in_para(p, PORTABILITY_OLD, PORTABILITY_NEW)
            done = True
            break
    report["Portability sentence simplified"] = "OK" if done else "SKIP"

    done = False
    for p in doc.paragraphs:
        if DA_OLD_SNIPPET in p.text:
            replace_in_para(p, DA_OLD, DA_NEW)
            done = True
            break
    report["Data availability updated"] = "OK" if done else "SKIP"

    done = False
    for p in doc.paragraphs:
        if AC_PLACEHOLDER in p.text:
            replace_in_para(p, AC_PLACEHOLDER, AC_FULL)
            done = True
            break
    report["CRediT author contributions"] = "OK" if done else "SKIP"

    done = False
    for p in doc.paragraphs:
        if GITHUB_PLACEHOLDER in p.text:
            replace_in_para(p, GITHUB_PLACEHOLDER, GITHUB_URL)
            done = True
            break
    if not done:
        report["GitHub URL"] = "Already present" if any(GITHUB_URL in p.text for p in doc.paragraphs) else "NOT FOUND"
    else:
        report["GitHub URL"] = "OK replaced"

    # Verifier para
    if any("claim verifier (Stage 5) operates" in p.text for p in doc.paragraphs):
        report["Verifier description"] = "Already present"
    else:
        done = False
        for p in doc.paragraphs:
            if SECTION_42_TEXT in p.text and p.style.name.lower().startswith("heading"):
                new_p = OxmlElement("w:p")
                p._p.addprevious(new_p)
                for q in doc.paragraphs:
                    if q._p is new_p:
                        q.add_run(VERIFIER_INTRO_PARA)
                        done = True
                        break
                break
        report["Verifier description"] = "OK inserted" if done else "NOT FOUND"

    return report


def remove_all_image_paragraphs(doc: Document) -> int:
    """Remove all paragraphs that contain images. Returns count removed."""
    to_remove = [p for p in doc.paragraphs if has_picture(p)]
    for p in to_remove:
        p._p.getparent().remove(p._p)
    return len(to_remove)


def insert_figure_before_caption(doc: Document, label: str, fig_path: Path) -> bool:
    """Insert a figure paragraph directly before the caption paragraph that starts with label."""
    for p in doc.paragraphs:
        if p.text.strip().startswith(label + " "):
            # Create image paragraph
            new_p = OxmlElement("w:p")
            p._p.addprevious(new_p)
            # Find the new paragraph object
            for q in doc.paragraphs:
                if q._p is new_p:
                    run = q.add_run()
                    run.add_picture(str(fig_path), width=Inches(5.5))
                    q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    return True
    return False


def main():
    print(f"Loading {SRC.name} ...")
    doc = Document(str(SRC))

    print("\nStep 1: Apply text edits ...")
    text_report = apply_text_edits(doc)
    for k, v in text_report.items():
        print(f"  {k}: {v}")

    print("\nStep 2: Remove all existing image paragraphs ...")
    n_removed = remove_all_image_paragraphs(doc)
    print(f"  Removed {n_removed} image paragraphs")

    print("\nStep 3: Re-insert figures directly before their captions ...")
    for label, fig_path in FIGURE_MAP.items():
        ok = insert_figure_before_caption(doc, label, fig_path)
        status = "OK" if ok else "CAPTION NOT FOUND"
        print(f"  {label}: {status}")

    doc.save(str(DST))
    print(f"\nSaved -> {DST.name}")

    # Final audit
    print("\n=== FINAL FIGURE/CAPTION ORDER AUDIT ===")
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        hp = has_picture(p)
        is_cap = any(txt.startswith(lbl + " ") for lbl in FIGURE_MAP)
        if hp or is_cap:
            tp = "IMAGE  " if hp else "CAPTION"
            prev_is_image = (i > 0 and has_picture(doc.paragraphs[i-1])) if is_cap else False
            check = " [CORRECTLY BEFORE CAPTION]" if (hp and i+1 < len(doc.paragraphs) and
                doc.paragraphs[i+1].text.strip()[:5] in [lbl[:5] for lbl in FIGURE_MAP]) else ""
            print(f"  Para {i:3d} [{tp}]: {txt[:70]}{check}")


if __name__ == "__main__":
    main()
